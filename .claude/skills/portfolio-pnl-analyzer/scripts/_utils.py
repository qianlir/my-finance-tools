#!/usr/bin/env python3
"""共享工具:代码规范化 + Excel schema 验证 + 投资人风格的小数格式化"""
import openpyxl


def normalize_code(c) -> str:
    """同花顺代码统一格式:
    - 5 位数字(港股 09988) → 保留 5 位
    - <=6 位数字 → zfill(6)
    - 其他原样返回
    """
    if c is None:
        return ''
    s = str(c).strip()
    if not s.isdigit():
        return s
    if len(s) == 5:
        return s
    if len(s) <= 6:
        return s.zfill(6)
    return s


REQUIRED_SHEETS = {
    '持仓数据': ['代码', '名称', '持有金额', '持有数量', '今年盈亏', '累计盈亏', '最新价'],
    '已清仓': ['清仓日期', '代码', '名称', '总盈亏'],
    '交易记录': ['成交日期', '代码', '名称', '交易类别', '成交数量', '成交价格', '发生金额', '费用'],
}


def validate_excel(xlsx_path: str) -> list:
    """检查 Excel 必需 sheet + 必需列。返回错误列表,空 = OK。"""
    errors = []
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as e:
        return [f"无法打开 Excel: {e}"]

    for sheet_name, required_cols in REQUIRED_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            errors.append(f"缺 sheet: {sheet_name}")
            continue
        ws = wb[sheet_name]
        # 读第一行表头
        headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
        missing = [col for col in required_cols if col not in headers]
        if missing:
            errors.append(f"sheet '{sheet_name}' 缺列: {missing}")
    wb.close()
    return errors


def fmt_money(v) -> str:
    if v is None: return 'N/A'
    return f'¥{v:,.0f}'


def fmt_pct(v, small_basis=False) -> str:
    """格式化百分比。small_basis=True 时返回 N/M(基数太小不显著)"""
    if v is None or v == 'N/A':
        return 'N/A'
    if isinstance(v, str):
        return v
    if small_basis:
        return 'N/M'  # Not Meaningful — 投资界惯例
    return f'{v:.2f}%'


# 投资界惯例:仓位/PL 基数小于此阈值时,收益率显示 N/M
SMALL_BASIS_THRESHOLD = 10000  # ¥1 万
