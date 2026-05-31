#!/usr/bin/env python3
"""LOF 申购套利分析 — 识别场外申购(fee≥0.02%) + 配对首次卖出日均价。

环境变量同上。

产出:
  data/lof_analysis.json — LOF 套利统计(给 05_report 用)
"""
import os
import csv
import json
import openpyxl
from collections import defaultdict
from datetime import datetime
from pathlib import Path

SRC = Path(os.environ['PORTFOLIO_XLSX']).expanduser()
WORK = Path(os.environ['PORTFOLIO_WORK_DIR']).expanduser()
TODAY = os.environ['PORTFOLIO_TODAY']
DATA = WORK / 'data'

FEE_THRESHOLD = 0.0002  # 0.02% 区分场内(≤0.01%) vs 申购(≥0.03%)


def is_lof(code):
    if not code.isdigit() or len(code) != 6:
        return False
    return code[:2] in ('16', '50') or code[:3] == '163'


def normalize_code(c):
    if c is None: return ''
    s = str(c).strip()
    if not s.isdigit(): return s
    if len(s) == 5: return s
    if len(s) <= 6: return s.zfill(6)
    return s


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

    # 1. 收集交易
    ws = wb['交易记录']
    all_trades = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        code = r[2]
        if not code: continue
        c = normalize_code(code)
        all_trades[c].append({
            'date': str(r[0])[:10] if r[0] else '',
            'name': r[3], 'type': r[4],
            'qty': float(r[5] or 0), 'amt': float(r[7] or 0), 'fee': float(r[9] or 0),
        })

    # 2. 同花顺持仓表
    ths = {}
    ws2 = wb['持仓数据']
    for r in ws2.iter_rows(min_row=2, values_only=True):
        c = r[0]
        if not c or str(c) == '汇总': continue
        c_s = normalize_code(c)
        ths[c_s] = {'name': r[1], 'ytd': float(r[15] or 0), 'cum': float(r[11] or 0),
                     'mv': float(r[2] or 0)}

    # 3. 已清仓 2026
    ws3 = wb['已清仓']
    cleared = defaultdict(float)
    for r in ws3.iter_rows(min_row=2, values_only=True):
        d = str(r[0])[:10] if r[0] else ''
        if d[:4] != TODAY[:4]: continue
        c = normalize_code(r[1])
        cleared[c] += float(r[3] or 0)
    wb.close()

    # 4. 申购套利配对
    lof_arb = []
    for code, trades in all_trades.items():
        if not is_lof(code): continue
        trades.sort(key=lambda x: x['date'])
        name = next((t['name'] for t in trades if t['name']), '')

        subs = []
        sells_by_date = defaultdict(list)
        for t in trades:
            if t['type'] in ('买入',) and t['qty'] > 0 and abs(t['amt']) > 10:
                if t['fee'] / abs(t['amt']) >= FEE_THRESHOLD:
                    subs.append(t)
            elif t['type'] in ('卖出', '卖券还款'):
                sells_by_date[t['date']].append(t)

        if not subs: continue

        # 卖出日均价
        sell_day_avg = {}
        for d, day_sells in sells_by_date.items():
            ta = sum(abs(s['amt']) for s in day_sells)
            tq = sum(s['qty'] for s in day_sells)
            tf = sum(s['fee'] for s in day_sells)
            if tq > 0:
                sell_day_avg[d] = (ta - tf) / tq
        sell_dates = sorted(sell_day_avg.keys())

        cycles = []
        for sub in subs:
            sc = (abs(sub['amt']) + sub['fee']) / sub['qty']
            for sd in sell_dates:
                if sd <= sub['date']: continue
                days = (datetime.strptime(sd, '%Y-%m-%d') - datetime.strptime(sub['date'], '%Y-%m-%d')).days
                if days >= 1:
                    sp = sell_day_avg[sd]
                    pl = (sp - sc) * sub['qty']
                    ret = (sp / sc - 1) * 100
                    cycles.append({'pl': round(pl, 0), 'ret': round(ret, 1),
                                   'days': days, 'win': pl > 0})
                    break

        if cycles:
            lof_arb.append({
                'code': code, 'name': name,
                'sub_count': len(subs), 'matched': len(cycles),
                'wins': sum(1 for c in cycles if c['win']),
                'total_pl': round(sum(c['pl'] for c in cycles), 0),
                'avg_days': round(sum(c['days'] for c in cycles) / len(cycles), 1),
            })

    # 5. LOF 盈亏排行(同花顺口径, 同名合并)
    name_groups = {
        '黄金': ['164701', '160719', '161116'],
        '油气': ['162411', '160723', '161129', '163208', '160416'],
        '互联网QD': ['160644'], '海外科技': ['501312'], '全球芯片': ['501225'],
        '商品': ['160216'], '标普科技': ['161128'], '美国消费': ['162415'],
        '印度基金': ['164824'], '白银基金': ['161226'], '南方香港': ['160125'],
        '美国REIT': ['160140'], '纳指LOF': ['161130'], '标普500': ['161125'],
        '标普医药': ['161126'], '标普生物': ['161127'], '美元债': ['501300'],
        '香港小盘': ['161124'], '广发石油': ['162719'],
    }
    assigned = set()
    for codes in name_groups.values():
        assigned.update(codes)

    # 找做过申购的标的
    sub_codes = {a['code'] for a in lof_arb}

    pnl_ranking = []
    for gname, codes in name_groups.items():
        items = [(c, ths[c]) for c in codes if c in ths and c in sub_codes]
        if not items: continue
        ytd = sum(d['ytd'] for _, d in items) + sum(cleared.get(c, 0) for c, _ in items)
        cum = sum(d['cum'] for _, d in items)
        mv = sum(d['mv'] for _, d in items)
        pnl_ranking.append({'name': gname, 'ytd': round(ytd, 0), 'cum': round(cum, 0), 'mv': round(mv, 0)})

    # 未归组
    for c in sub_codes:
        if c in assigned or c not in ths: continue
        d = ths[c]
        ytd = d['ytd'] + cleared.get(c, 0)
        pnl_ranking.append({'name': d['name'], 'ytd': round(ytd, 0), 'cum': round(d['cum'], 0), 'mv': round(d['mv'], 0)})

    pnl_ranking.sort(key=lambda x: -x['ytd'])

    # 6. 汇总
    total_matched = sum(a['matched'] for a in lof_arb)
    total_wins = sum(a['wins'] for a in lof_arb)
    total_arb_pl = sum(a['total_pl'] for a in lof_arb)
    total_ytd = sum(p['ytd'] for p in pnl_ranking)
    total_cum = sum(p['cum'] for p in pnl_ranking)

    result = {
        'arb_summary': {
            'total_codes': len(lof_arb),
            'total_matched': total_matched,
            'total_wins': total_wins,
            'win_rate': round(total_wins / total_matched * 100, 0) if total_matched else 0,
            'total_arb_pl': total_arb_pl,
        },
        'arb_by_code': sorted(lof_arb, key=lambda x: -x['total_pl']),
        'pnl_ranking': pnl_ranking,
        'pnl_summary': {
            'total_ytd': total_ytd,
            'total_cum': total_cum,
            'product_count': len(pnl_ranking),
            'code_count': len(sub_codes),
        },
    }

    with open(DATA / 'lof_analysis.json', 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f'LOF 套利: {len(lof_arb)} 个标的, {total_matched} 笔配对, 成功率 {result["arb_summary"]["win_rate"]:.0f}%, P&L ¥{total_arb_pl:+,.0f}')
    print(f'LOF 盈亏: {len(pnl_ranking)} 个品种, YTD ¥{total_ytd:+,.0f}, 累计 ¥{total_cum:+,.0f}')
    print(f'✓ {DATA / "lof_analysis.json"}')


if __name__ == '__main__':
    main()
